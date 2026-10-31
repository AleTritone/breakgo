# SEZIONE: Importazione ed esportazione Excel
import io
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .schema import TABLES, LABELS

NAVY = '14283F'
TEAL = '16B8A6'

def workbook(tables):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        for name, df in tables.items():
            df.to_excel(writer, sheet_name=LABELS.get(name,name)[:31], index=False)
            ws = writer.sheets[LABELS.get(name,name)[:31]]
            ws.freeze_panes = 'A2'
            ws.sheet_view.showGridLines = False
            ws.row_dimensions[1].height = 29
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor=NAVY)
                cell.font = Font(name='Aptos', size=11, bold=True, color='FFFFFF')
                cell.alignment = Alignment(vertical='center')
            for j, col in enumerate(df.columns, 1):
                lengths = [len(str(col))] + [len(str(x)) for x in df[col].head(80).fillna('')]
                ws.column_dimensions[get_column_letter(j)].width = min(47, max(16, max(lengths)+3))
            ws.auto_filter.ref = ws.dimensions
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name='Aptos', size=10, color='147D76')
    return output.getvalue()

def read_upload(upload, table):
    content = upload.getvalue()
    sheets = pd.read_excel(io.BytesIO(content), sheet_name=None, dtype=str, keep_default_na=False)
    expected = TABLES[table]
    matching = [df for df in sheets.values() if set(expected).issubset(df.columns)]
    if not matching:
        raise ValueError('Nessun foglio contiene le colonne richieste: ' + ', '.join(expected))
    return matching[0][expected].copy()
